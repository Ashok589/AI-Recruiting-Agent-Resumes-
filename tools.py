
import re
from pathlib import Path

import pandas as pd
from docx import Document
from pypdf import PdfReader


# Common aliases for equivalent skill names.
SKILL_ALIASES = {
    "powerbi": "power bi",
    "ms excel": "excel",
    "microsoft excel": "excel",
    "postgres": "postgresql",
    "scikit learn": "scikit-learn",
    "sklearn": "scikit-learn",
    "ml": "machine learning",
    "genai": "generative ai",
    "gen ai": "generative ai",
    "llm": "llms",
    "rest api": "rest apis",
    "js": "javascript",
}


def normalize_skill(value):
    """Normalize skill names for consistent comparison."""
    value = str(value or "").strip().lower()
    value = re.sub(r"[^a-z0-9+#.\s-]", " ", value)
    value = re.sub(r"\s+", " ", value).strip()

    return SKILL_ALIASES.get(value, value)


def extract_resume_text(uploaded_file):
    """Tool 1: Extract text from PDF, DOCX, or TXT resumes."""
    filename = str(getattr(uploaded_file, "name", "")).lower()

    try:
        if filename.endswith(".pdf"):
            reader = PdfReader(uploaded_file)
            text = "\n".join(
                page.extract_text() or ""
                for page in reader.pages
            )

        elif filename.endswith(".docx"):
            document = Document(uploaded_file)

            paragraphs = [
                paragraph.text
                for paragraph in document.paragraphs
                if paragraph.text.strip()
            ]

            # Include tables because many resumes use them for layout.
            for table in document.tables:
                for row in table.rows:
                    paragraphs.append(
                        " | ".join(
                            cell.text.strip()
                            for cell in row.cells
                        )
                    )

            text = "\n".join(paragraphs)

        elif filename.endswith(".txt"):
            if hasattr(uploaded_file, "getvalue"):
                file_bytes = uploaded_file.getvalue()
            else:
                file_bytes = uploaded_file.read()

            text = file_bytes.decode(
                "utf-8",
                errors="replace",
            )

        else:
            raise ValueError(
                "Unsupported file. Upload PDF, DOCX, or TXT."
            )

        text = text.strip()

        if not text:
            raise ValueError(
                "No readable text found. The PDF may be scanned; "
                "OCR may be required."
            )

        return text

    except ValueError:
        raise
    except Exception as exc:
        raise ValueError(
            f"Resume extraction failed: {exc}"
        ) from exc


def search_candidate_database(
    query,
    candidates=None,
    limit=5,
):
    """
    Tool 2: Search internal candidate records using query terms.

    candidates can be a list of dictionaries or records loaded from CSV.
    """
    if not query or not str(query).strip():
        return []

    if candidates is None:
        candidates = []

    terms = set(
        re.findall(
            r"[a-zA-Z0-9+#.-]+",
            str(query).lower(),
        )
    )

    # Avoid matching very common words.
    stop_words = {
        "and", "the", "for", "with", "from",
        "this", "that", "are", "was", "has",
    }
    terms = {
        term for term in terms
        if len(term) > 1 and term not in stop_words
    }

    results = []

    for candidate in candidates:
        if not isinstance(candidate, dict):
            continue

        def field_text(field):
            value = candidate.get(field, "")
            if value is None:
                return ""
            if isinstance(value, list):
                return " ".join(str(item) for item in value)
            return str(value)

        searchable = " ".join(
            field_text(field)
            for field in (
                "name",
                "skills",
                "experience",
                "role",
                "summary",
            )
        ).lower()

        matched_terms = sorted(
            term for term in terms
            if re.search(
                rf"(?<!\w){re.escape(term)}(?!\w)",
                searchable,
            )
        )

        if matched_terms:
            results.append({
                **candidate,
                "matched_terms": matched_terms,
                "search_match_count": len(matched_terms),
            })

    results.sort(
        key=lambda item: item["search_match_count"],
        reverse=True,
    )

    return results[:max(0, limit)]


def score_candidate(candidate_skills, required_skills):
    """
    Tool 3: Calculate a transparent required-skill match score.

    The score is:
        matched required skills / total required skills * 100

    This measures skill overlap, not the probability of job success.
    """
    candidate_skills = candidate_skills or []
    required_skills = required_skills or []

    if isinstance(candidate_skills, str):
        candidate_skills = [candidate_skills]

    if isinstance(required_skills, str):
        required_skills = [required_skills]

    # Normalize candidate skills while preserving readable labels.
    normalized_candidates = {}

    for skill in candidate_skills:
        original = str(skill or "").strip()
        normalized = normalize_skill(original)

        if normalized:
            normalized_candidates.setdefault(
                normalized,
                original,
            )

    # Deduplicate required skills after normalization.
    normalized_required = {}

    for skill in required_skills:
        original = str(skill or "").strip()
        normalized = normalize_skill(original)

        if normalized:
            normalized_required.setdefault(
                normalized,
                original,
            )

    if not normalized_required:
        return {
            "score": None,
            "matched_count": 0,
            "required_count": 0,
            "matched_skills": [],
            "missing_skills": [],
            "message": (
                "No required skills were identified. "
                "A match score cannot be calculated."
            ),
        }

    matched_skills = []
    missing_skills = []

    for normalized, required_label in normalized_required.items():
        if normalized in normalized_candidates:
            matched_skills.append({
                "required_skill": required_label,
                "resume_skill": normalized_candidates[normalized],
            })
        else:
            missing_skills.append(required_label)

    score = round(
        len(matched_skills)
        / len(normalized_required)
        * 100,
        1,
    )

    return {
        "score": score,
        "matched_count": len(matched_skills),
        "required_count": len(normalized_required),
        "matched_skills": matched_skills,
        "missing_skills": missing_skills,
        "message": (
            f"{len(matched_skills)} of "
            f"{len(normalized_required)} required skills matched."
        ),
    }


def save_candidate_record(
    record,
    csv_path="candidates.csv",
):
    """
    Tool 4: Save a candidate record to the local CSV database.

    Records with the same name and email are updated instead of duplicated.
    """
    if not isinstance(record, dict):
        raise ValueError(
            "Candidate record must be a dictionary."
        )

    path = Path(csv_path)
    path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    new_record = pd.DataFrame([record])

    if path.exists():
        try:
            existing = pd.read_csv(path).fillna("")
        except pd.errors.EmptyDataError:
            existing = pd.DataFrame()

        combined = pd.concat(
            [existing, new_record],
            ignore_index=True,
        )

        # Use available identity columns for duplicate detection.
        identity_columns = [
            column
            for column in ("name", "email")
            if column in combined.columns
        ]

        if identity_columns:
            combined = combined.drop_duplicates(
                subset=identity_columns,
                keep="last",
            )
    else:
        combined = new_record

    combined.to_csv(
        path,
        index=False,
        encoding="utf-8",
    )

    return str(path)
